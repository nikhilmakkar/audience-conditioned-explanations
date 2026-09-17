"""Create a human-friendly Excel workbook for prospective stimulus review."""

from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "results" / "prospective_stimulus_review.csv"
OUTPUT = ROOT / "results" / "PROSPECTIVE_STIMULUS_REVIEW.xlsx"


def main() -> None:
    with SOURCE.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"

    headers = [
        "Method",
        "Source passage",
        "Technical summary",
        "Accessible summary",
        "Passage internally plausible?",
        "Technical summary accurate?",
        "Accessible summary accurate?",
        "Both summaries comparably useful?",
        "Notes",
    ]
    ws.append(headers)

    source_keys = [
        "source",
        "passage",
        "technical_summary",
        "accessible_summary",
        "passage_accurate",
        "technical_summary_accurate",
        "accessible_summary_accurate",
        "pair_comparably_useful",
        "notes",
    ]
    for row in rows:
        ws.append([row[key] for key in source_keys])

    navy = "17365D"
    pale_blue = "DDEBF7"
    pale_yellow = "FFF2CC"
    grey = "E7E6E6"
    white = "FFFFFF"
    thin = Side(style="thin", color="B7B7B7")

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 145
        for col_idx in range(1, 10):
            cell = ws.cell(row_idx, col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if col_idx <= 4:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
            else:
                cell.fill = PatternFill("solid", fgColor=pale_yellow)
        ws.cell(row_idx, 1).font = Font(bold=True)
        for col_idx in range(5, 9):
            ws.cell(row_idx, col_idx).alignment = Alignment(
                horizontal="center", vertical="center"
            )

    widths = {
        "A": 16,
        "B": 62,
        "C": 48,
        "D": 48,
        "E": 21,
        "F": 21,
        "G": 21,
        "H": 24,
        "I": 38,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    ws.sheet_view.showGridLines = False

    validation = DataValidation(type="list", formula1='"1,0"', allow_blank=True)
    validation.promptTitle = "Enter 1 or 0"
    validation.prompt = "1 = yes; 0 = no"
    validation.errorTitle = "Use 1 or 0"
    validation.error = "Please choose 1 (yes) or 0 (no)."
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(f"E2:H{ws.max_row}")

    guide = wb.create_sheet("Instructions")
    instructions = [
        ("Prospective stimulus review", True),
        ("Blue cells are the material to read. Yellow cells are for your ratings.", False),
        ("For each method, compare the source passage with both candidate summaries.", False),
        ("Enter 1 for yes and 0 for no in each of the four rating columns.", False),
        ("Passage internally plausible?: Is the description coherent and broadly plausible?", False),
        ("Technical summary accurate?: Does it faithfully preserve the source passage?", False),
        ("Accessible summary accurate?: Does it faithfully preserve the source passage without requiring field-specific vocabulary?", False),
        ("Both summaries comparably useful?: Is terminology—not correctness, detail, or quality—the main difference?", False),
        ("Use Notes for any 0 or uncertainty. You do not need to research all eight fields.", False),
    ]
    for text, heading in instructions:
        guide.append([text])
        guide.cell(guide.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        if heading:
            guide.cell(guide.max_row, 1).font = Font(size=16, bold=True, color=navy)
    guide.column_dimensions["A"].width = 105
    guide.row_dimensions[1].height = 28
    guide.sheet_view.showGridLines = False

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
